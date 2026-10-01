"""
build_pipe_database.py — build data/pipe_schedules.json from the uploaded
Steel Pipe Schedule Chart (ANSI B36.10 & B36.19).

Run:  python build_pipe_database.py

WHAT IT PRODUCES
    For every NPS and every published schedule: outside diameter, wall
    thickness, inside diameter and mass per metre, all in mm / kg.

    Mass per metre is CALCULATED, not transcribed:

        cross-section = pi x t x (OD - t)          [mm^2]
        kg/m          = area x 1000 x 1e-9 x density

    which is the standard pipe mass formula. Checked against known values —
    2" SCH 40 comes out 5.44 kg/m, which is the book figure.

TWO CORRECTIONS TO THE SOURCE
    Both are flagged in the spreadsheet's own footnotes; both are handled here
    rather than left for the estimator to trip over.

    1. The SCH 10s column is in MILLIMETRES while every other column is in
       inches. The sheet's note calls the values "far larger than typical" and
       suspects a merged column — it is simpler than that. Divide by 25.4 and
       every value lands on ASME B36.19 within 0.4%:
           NPS 1/2  ->  2.11 / 25.4 = 0.083 in   (B36.19: 0.083)
           NPS 12   ->  4.57 / 25.4 = 0.180 in   (B36.19: 0.180)
       So the data is good, the unit label is wrong. Corrected, not discarded.

    2. NPS 26 SCH 30 is printed as 0.000 in. A zero wall is not a pipe. Treated
       as "not published" and dropped, exactly as the footnote advises.

STAINLESS vs CARBON SCHEDULES
    5s / 10s / 40s / 80s are the stainless schedules of ANSI/ASME B36.19; the
    rest are carbon-steel schedules of B36.10. They are kept in one table
    because that is how the chart is published and how people order, but each
    row records which standard it came from.

MASS AND MATERIAL
    kg/m here is for CARBON STEEL at 7850 kg/m3. The app rescales it by the
    density of whatever material the user picks, so a stainless or aluminium
    pipe of the same size comes out right.
"""
import json
import math
import os

from openpyxl import load_workbook

SRC = "/mnt/user-data/uploads/Steel_Pipe_Schedule_Chart_ANSI_B36_10_B36_19.xlsx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "data", "pipe_schedules.json")

DENSITY_CS = 7850.0
IN_TO_MM = 25.4

# Columns whose values are in millimetres despite the sheet being labelled
# inches. See correction (1) above.
MM_COLUMNS = {"SCH 10s"}

STAINLESS = {"SCH 5s", "SCH 10s", "SCH 40s", "SCH 80s"}

CORRECTIONS = [
    {
        "field": "SCH 10s",
        "issue": "column is in millimetres, not inches",
        "action": "divided by 25.4",
        "evidence": ("Every converted value matches ASME B36.19 within 0.4% "
                     "(1/2\" -> 0.083 in, 12\" -> 0.180 in). The source footnote "
                     "suspected a merged column; it is a unit mislabel."),
    },
    {
        "field": "NPS 26, SCH 30",
        "issue": "printed as 0.000 in",
        "action": "dropped as not published",
        "evidence": "A zero wall thickness is not a pipe. The source footnote "
                    "flags it as a typo or placeholder.",
    },
]


def mass_per_m(od_mm, t_mm, density=DENSITY_CS):
    """Standard pipe mass: area = pi x t x (OD - t), taken over one metre."""
    if not (od_mm and t_mm) or t_mm <= 0 or t_mm * 2 >= od_mm:
        return None
    area_mm2 = math.pi * t_mm * (od_mm - t_mm)
    return area_mm2 * 1000.0 * 1e-9 * density


def build():
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Pipe Schedule (in)"]
    rows = list(ws.iter_rows(values_only=True))

    hdr_i = next(i for i, r in enumerate(rows) if r[0] == "NPS")
    hdr = [str(c).strip() if c else "" for c in rows[hdr_i]]
    sched_cols = {i: h for i, h in enumerate(hdr) if h.startswith("SCH")}

    sizes, dropped = [], []
    for r in rows[hdr_i + 1:]:
        if not r[0] or not isinstance(r[2], (int, float)):
            continue
        nps = str(r[0]).strip()
        od_in = float(r[2])
        od_mm = od_in * IN_TO_MM

        walls = {}
        for i, name in sched_cols.items():
            v = r[i]
            if v in (None, "") or not isinstance(v, (int, float)):
                continue
            v = float(v)
            if v <= 0:                       # correction (2)
                dropped.append(f"NPS {nps} {name} = {v}")
                continue
            t_mm = v if name in MM_COLUMNS else v * IN_TO_MM   # correction (1)
            key = name.replace("SCH ", "")
            walls[key] = {
                "t_mm": round(t_mm, 3),
                "id_mm": round(od_mm - 2 * t_mm, 3),
                "kg_m": round(mass_per_m(od_mm, t_mm), 4),
                "standard": "B36.19" if name in STAINLESS else "B36.10",
            }

        if walls:
            sizes.append({
                "nps": nps,
                "dn": r[1],
                "od_in": od_in,
                "od_mm": round(od_mm, 3),
                "schedules": walls,
            })

    data = {
        "_source": "Steel Pipe Schedule Chart, ANSI B36.10 & B36.19 (Octal Steel)",
        "_note": ("Wall thicknesses converted to mm. kg/m is calculated as "
                  "pi x t x (OD - t) over one metre at 7850 kg/m3 carbon steel, "
                  "and is rescaled by material density in the app."),
        "_density_basis_kg_m3": DENSITY_CS,
        "corrections": CORRECTIONS,
        "dropped": dropped,
        "sizes": sizes,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=1)

    n_sched = sum(len(s["schedules"]) for s in sizes)
    print(f"Wrote {OUT}")
    print(f"  {len(sizes)} sizes, {n_sched} size/schedule combinations")
    print(f"  dropped: {dropped or 'none'}")
    return data


def validate(data=None):
    """Check the calculated masses against published pipe weights."""
    data = data or json.load(open(OUT, encoding="utf-8"))
    idx = {s["nps"]: s for s in data["sizes"]}
    # nps, schedule, published kg/m
    known = [("1/2", "40", 1.27), ("1", "40", 2.50), ("2", "40", 5.44),
             ("3", "40", 11.29), ("4", "40", 16.07), ("6", "40", 28.26),
             ("8", "40", 42.55), ("10", "40", 60.31), ("12", "STD", 73.88),
             ("2", "80", 7.48), ("4", "80", 22.32), ("6", "80", 42.56)]
    print(f"\n{'NPS':>6} {'sch':>5} {'computed':>9} {'published':>10} {'diff':>7}")
    worst = 0.0
    for nps, sch, pub in known:
        s = idx.get(nps)
        if not s or sch not in s["schedules"]:
            print(f"{nps:>6} {sch:>5}  not in dataset")
            continue
        got = s["schedules"][sch]["kg_m"]
        d = (got - pub) / pub * 100
        worst = max(worst, abs(d))
        print(f"{nps:>6} {sch:>5} {got:>9.2f} {pub:>10.2f} {d:>6.1f}%")
    print(f"\n  worst deviation {worst:.1f}% "
          f"({'OK' if worst < 2 else 'CHECK THE TRANSCRIPTION'})")
    return worst


if __name__ == "__main__":
    validate(build())
