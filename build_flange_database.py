import json
from pathlib import Path
import openpyxl


# ============================================================
# FILE LOCATIONS
# ============================================================

BASE_DIR = Path(__file__).parent

EXCEL_FILE = (
    BASE_DIR /
    "ANSI_B16.5_Forged_Flanges_Datasheet.xlsx"
)

OUTPUT_FILE = (
    BASE_DIR /
    "flange_database.json"
)


# ============================================================
# PRESSURE CLASS SHEETS
# ============================================================

CLASS_SHEETS = {
    150: "Class150",
    300: "Class300",
    400: "Class400",
    600: "Class600",
    900: "Class900",
    1500: "Class1500",
    2500: "Class2500",
}


# ============================================================
# FLANGE TYPES
# ============================================================

FLANGE_TYPES = [
    "Welding Neck",
    "Slip-on",
    "Socket Welding",
    "Threaded",
    "Lap Joint",
    "Blind",
]


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean(value):
    """Clean Excel cell values."""

    if value is None:
        return None

    if isinstance(value, str):

        value = value.strip()

        if value == "":
            return None

    return value


def number(value):
    """
    Convert numeric Excel values to float.
    Keep text such as 'See purchaser' unchanged.
    """

    value = clean(value)

    if value is None:
        return None

    try:
        return float(value)

    except (ValueError, TypeError):

        return value


def normalize(text):
    """
    Normalize a header so that small differences in
    wording/capitalization do not cause problems.
    """

    if text is None:
        return ""

    return (
        str(text)
        .strip()
        .lower()
        .replace("\n", " ")
        .replace("  ", " ")
    )


# ============================================================
# FIND TABLE HEADER
# ============================================================

def find_header_row(ws, table_name):
    """
    Find the row containing 'Nominal Pipe Size'.

    table_name is only used for error messages.
    """

    for row_number in range(1, ws.max_row + 1):

        value = ws.cell(
            row_number,
            1
        ).value

        if value is None:
            continue

        text = str(value).strip().lower()

        if text == "nominal pipe size":

            return row_number

    raise ValueError(
        f"Could not find header for {table_name}"
    )


# ============================================================
# READ TABLE INTO DICTIONARY
# ============================================================

def read_table(ws, header_row):
    """
    Convert one Excel table into:

        NPS -> column-name -> value
    """

    headers = {}

    for column in range(
        1,
        ws.max_column + 1
    ):

        value = ws.cell(
            header_row,
            column
        ).value

        if value is not None:

            headers[
                normalize(value)
            ] = column

    records = {}

    row_number = header_row + 1

    while row_number <= ws.max_row:

        nps = clean(
            ws.cell(
                row_number,
                1
            ).value
        )

        # Stop when we reach a blank row
        # or the Notes section.
        if nps is None:
            break

        if str(nps).strip().lower() == "notes:":
            break

        row_data = {}

        for header, column in headers.items():

            row_data[header] = clean(
                ws.cell(
                    row_number,
                    column
                ).value
            )

        records[str(nps)] = row_data

        row_number += 1

    return headers, records


# ============================================================
# GET VALUE FROM RECORD
# ============================================================

def get_value(record, possible_headers):
    """
    Try several possible Excel header names.

    This is useful because Class 400/900 etc.
    use slightly different wording.
    """

    for header in possible_headers:

        header = normalize(header)

        if header in record:

            return number(
                record[header]
            )

    return None


# ============================================================
# BUILD ONE FLANGE RECORD
# ============================================================

def build_flange_record(
    pressure_class,
    nps,
    dimensions,
    bolting
):

    flange_records = []

    # --------------------------------------------------------
    # Common dimensions
    # --------------------------------------------------------

    common_dimensions = {

        "outside_d": get_value(
            dimensions,
            [
                "Outside Diam. D"
            ]
        ),

        "hub_base_x": get_value(
            dimensions,
            [
                "Diam. at Base of Hub X"
            ]
        ),

        "raised_face_od_g": get_value(
            dimensions,
            [
                "O.D. of Raised Face G"
            ]
        ),

        "thickness_t": get_value(
            dimensions,
            [
                "Thickness t"
            ]
        ),

        "hub_bevel_diameter_a": get_value(
            dimensions,
            [
                "Diam. of Hub at Bevel A"
            ]
        ),

        "fillet_radius_r": get_value(
            dimensions,
            [
                "Radius of Fillet R"
            ]
        ),

        "thread_length_q": get_value(
            dimensions,
            [
                "Thread Length Q"
            ]
        ),
    }


    # --------------------------------------------------------
    # Common bolting information
    # --------------------------------------------------------

    bolting_data = {

        "bolt_circle_diameter": get_value(
            bolting,
            [
                "Bolt Circle Diam"
            ]
        ),

        "number_of_holes": get_value(
            bolting,
            [
                "Number of Holes"
            ]
        ),

        "hole_diameter": get_value(
            bolting,
            [
                "Diam of Holes"
            ]
        ),

        "bolt_diameter_in": get_value(
            bolting,
            [
                "Diam of Bolts (in)"
            ]
        ),

        "socket_depth_y": get_value(
            bolting,
            [
                "Depth of Socket Y"
            ]
        ),

        "machine_bolt_length_raised_face":
            get_value(
                bolting,
                [
                    "Machine Bolt Length - Raised Face"
                ]
            ),

        "stud_bolt_length_raised_face":
            get_value(
                bolting,
                [
                    "Stud Bolt Length - Raised Face"
                ]
            ),

        "stud_bolt_length_ring_joint":
            get_value(
                bolting,
                [
                    "Stud Bolt Length - Ring Joint"
                ]
            ),

        "stud_bolt_length_025_raised_face":
            get_value(
                bolting,
                [
                    "Stud Bolt Length - 0.25in Raised Face"
                ]
            ),

        "stud_bolt_length_male_female":
            get_value(
                bolting,
                [
                    "Stud Bolt Length - Male-Female Tongue-Groove"
                ]
            ),
    }


    # --------------------------------------------------------
    # Weight columns
    # --------------------------------------------------------

    weight_map = {

        "Welding Neck": [
            "Wt. Welding Neck (kg)"
        ],

        "Slip-on": [
            "Wt. Slip-on & Threaded (kg)"
        ],

        "Threaded": [
            "Wt. Slip-on & Threaded (kg)"
        ],

        "Lap Joint": [
            "Wt. Lap Joint (kg)"
        ],

        "Blind": [
            "Wt. Blind (kg)"
        ],

        "Socket Welding": [
            "Wt. Socket Welding (kg)"
        ],
    }


    # --------------------------------------------------------
    # Bore columns
    # --------------------------------------------------------

    bore_map = {

        "Welding Neck": [
            "Bore - WN/Socket Welding B1",
            "Bore - Welding Neck B1"
        ],

        "Socket Welding": [
            "Bore - WN/Socket Welding B1",
            "Bore - Welding Neck B1"
        ],

        "Slip-on": [
            "Bore - Slip-on Socket Welding B2",
            "Bore - Slip-on B2"
        ],

        "Threaded": [
            "Counter Bore Min./Threaded Min. B"
        ],

        "Lap Joint": [
            "Bore - Lap Joint B3"
        ],

        "Blind": []
    }


    # --------------------------------------------------------
    # Hub length
    # --------------------------------------------------------

    hub_length_map = {

        "Welding Neck": [
            "Length thru Hub - Welding Neck T1"
        ],

        "Slip-on": [
            "Length thru Hub - Slip-on/Threaded/Socket Welding T2"
        ],

        "Threaded": [
            "Length thru Hub - Slip-on/Threaded/Socket Welding T2"
        ],

        "Socket Welding": [
            "Length thru Hub - Slip-on/Threaded/Socket Welding T2"
        ],

        "Lap Joint": [
            "Length thru Hub - Lap Joint T3"
        ],

        "Blind": []
    }


    # --------------------------------------------------------
    # Create records for every available flange type
    # --------------------------------------------------------

    for flange_type in FLANGE_TYPES:

        weight = get_value(
            bolting,
            weight_map[flange_type]
        )

        # If there is no weight value, that flange
        # type is not available for this size/class.
        if weight is None:
            continue

        bore = get_value(
            dimensions,
            bore_map[flange_type]
        )

        hub_length = get_value(
            dimensions,
            hub_length_map[flange_type]
        )

        dimensions_data = dict(
            common_dimensions
        )

        dimensions_data["bore"] = bore

        dimensions_data["hub_length"] = hub_length


        record = {

            "id": (
                f"ANSI-B16.5-"
                f"{pressure_class}-"
                f"{nps}-"
                f"{flange_type}"
            ),

            "standard": "ANSI B16.5",

            "pressure_class": pressure_class,

            "nps": str(nps),

            "flange_type": flange_type,

            "dimensions": dimensions_data,

            "bolting": dict(
                bolting_data
            ),

            "approximate_weight_kg": weight,

        }

        flange_records.append(record)

    return flange_records


# ============================================================
# BUILD COMPLETE DATABASE
# ============================================================

def build_database():

    if not EXCEL_FILE.exists():

        raise FileNotFoundError(
            f"\nExcel file not found:\n"
            f"{EXCEL_FILE}\n\n"
            f"Put the Excel file in the same folder "
            f"as this Python script."
        )


    print()
    print("=" * 60)
    print("BUILDING FLANGE DATABASE")
    print("=" * 60)
    print()


    workbook = openpyxl.load_workbook(
        EXCEL_FILE,
        data_only=True
    )


    database = {

        "database_name":
            "ANSI B16.5 Forged Flanges",

        "standard":
            "ANSI B16.5",

        "source_file":
            EXCEL_FILE.name,

        "units": {

            "dimensions":
                "mm",

            "weight":
                "kg",

            "bolt_diameter":
                "inch"

        },

        "flanges": []

    }


    # --------------------------------------------------------
    # Process each pressure class
    # --------------------------------------------------------

    for pressure_class, sheet_name in CLASS_SHEETS.items():

        print(
            f"Processing Class {pressure_class}..."
        )


        if sheet_name not in workbook.sheetnames:

            print(
                f"  WARNING: {sheet_name} not found"
            )

            continue


        ws = workbook[sheet_name]


        # Find Table A
        table_a_header = find_header_row(
            ws,
            "Table A"
        )


        # Find Table B
        # Search AFTER Table A
        table_b_header = None

        for row_number in range(
            table_a_header + 1,
            ws.max_row + 1
        ):

            value = ws.cell(
                row_number,
                1
            ).value

            if value is None:
                continue

            if (
                str(value)
                .strip()
                .lower()
                == "nominal pipe size"
            ):

                table_b_header = row_number

                break


        if table_b_header is None:

            raise ValueError(
                f"Could not find Table B "
                f"in {sheet_name}"
            )


        # Read both tables
        _, table_a = read_table(
            ws,
            table_a_header
        )

        _, table_b = read_table(
            ws,
            table_b_header
        )


        class_count = 0


        # ----------------------------------------------------
        # Match Table A and Table B using NPS
        # ----------------------------------------------------

        for nps, dimensions in table_a.items():

            if nps not in table_b:

                print(
                    f"  WARNING: "
                    f"NPS {nps} missing from Table B"
                )

                continue


            bolting = table_b[nps]


            records = build_flange_record(
                pressure_class,
                nps,
                dimensions,
                bolting
            )


            database["flanges"].extend(
                records
            )


            class_count += len(records)


        print(
            f"  → {class_count} flange records"
        )


    # --------------------------------------------------------
    # Save database
    # --------------------------------------------------------

    with open(
        OUTPUT_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            database,
            file,
            indent=4,
            ensure_ascii=False
        )


    print()
    print("=" * 60)
    print("DATABASE CREATED SUCCESSFULLY")
    print("=" * 60)
    print()

    print(
        "Total flange records:",
        len(database["flanges"])
    )

    print()

    print(
        "Database saved at:"
    )

    print(
        OUTPUT_FILE
    )

    print()


# ============================================================
# PROGRAM START
# ============================================================

if __name__ == "__main__":

    build_database()